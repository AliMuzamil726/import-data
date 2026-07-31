"""CSV, Excel and PDF exporters shared by every module."""
from __future__ import annotations

import csv
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .models import Report

BRAND = colors.HexColor("#1E5631")
BRAND_LIGHT = colors.HexColor("#E8F1EA")


def resolve(obj, attr: str):
    """Read a model attribute, property, callable or `get_x_display` value."""
    display = getattr(obj, f"get_{attr}_display", None)
    if callable(display):
        return display()
    value = getattr(obj, attr, "")
    if callable(value):
        value = value()
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, datetime):
        return timezone.localtime(value).strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    return value


def _log(request, kind: str, fmt: str, count: int) -> None:
    Report.objects.create(
        kind=kind, fmt=fmt, row_count=count,
        filters={k: v for k, v in request.GET.items() if v},
        generated_by=request.user if request.user.is_authenticated else None,
    )


def export_csv(request, rows, columns, *, title: str, kind: str, filename: str) -> HttpResponse:
    response = HttpResponse(content_type="text/csv")
    response["Content-Disposition"] = f'attachment; filename="{filename}.csv"'
    writer = csv.writer(response)
    writer.writerow([label for _, label in columns])
    count = 0
    for obj in rows:
        writer.writerow([resolve(obj, attr) for attr, _ in columns])
        count += 1
    _log(request, kind, "csv", count)
    return response


def export_xlsx(request, rows, columns, *, title: str, kind: str, filename: str) -> HttpResponse:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title[:31]

    header_fill = PatternFill("solid", fgColor="1E5631")
    header_font = Font(color="FFFFFF", bold=True, name="Calibri")
    for index, (_, label) in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="left", vertical="center")

    count = 0
    for row_index, obj in enumerate(rows, start=2):
        for col_index, (attr, _) in enumerate(columns, start=1):
            sheet.cell(row=row_index, column=col_index, value=resolve(obj, attr))
        count += 1

    for index, (_, label) in enumerate(columns, start=1):
        width = max(len(label) + 4, 14)
        sheet.column_dimensions[get_column_letter(index)].width = min(width, 42)
    sheet.freeze_panes = "A2"

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
    _log(request, kind, "xlsx", count)
    return response


def export_pdf(request, rows, columns, *, title: str, kind: str, filename: str) -> HttpResponse:
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        leftMargin=14 * mm, rightMargin=14 * mm, topMargin=14 * mm, bottomMargin=14 * mm,
        title=title, author="SAWIE",
    )
    styles = getSampleStyleSheet()
    heading = ParagraphStyle(
        "SawieHeading", parent=styles["Heading1"], textColor=BRAND, fontSize=18, spaceAfter=2,
    )
    meta = ParagraphStyle("SawieMeta", parent=styles["Normal"], fontSize=8,
                          textColor=colors.HexColor("#64748B"))
    cell = ParagraphStyle("SawieCell", parent=styles["Normal"], fontSize=7.5, leading=9.5)

    story = [
        Paragraph("SAWIE — Agricultural Intelligence Platform", meta),
        Paragraph(title, heading),
        Paragraph(
            f"Generated {timezone.localtime():%d %b %Y %H:%M} by "
            f"{getattr(request.user, 'display_name', 'system')}",
            meta,
        ),
        Spacer(1, 8),
    ]

    data = [[Paragraph(f"<b>{label}</b>", cell) for _, label in columns]]
    count = 0
    for obj in rows:
        data.append([Paragraph(str(resolve(obj, attr)), cell) for attr, _ in columns])
        count += 1

    table = Table(data, repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), BRAND),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, BRAND_LIGHT]),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(table)
    story.append(Spacer(1, 10))
    story.append(Paragraph(f"{count} rows", meta))

    doc.build(story)
    buffer.seek(0)
    response = HttpResponse(buffer.read(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename}.pdf"'
    _log(request, kind, "pdf", count)
    return response
